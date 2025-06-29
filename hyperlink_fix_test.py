#!/usr/bin/env python3
"""
Hyperlink Fix Test

This script creates a minimal test to reproduce and fix the first skill hyperlink issue.
We'll test the exact _add_cell_content_with_hyperlinks logic to identify the root cause.
"""

import sys
import os
from pathlib import Path

# Add src directory to path
sys.path.insert(0, str(Path(__file__).parent / 'src'))

try:
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
    
    # Import the document styles
    from skill_similarity_engine.webapp.career_analysis.config.document_styles import DocumentStyles
    
    DOCX_AVAILABLE = True
except ImportError as e:
    print(f"⚠️ Warning: python-docx not available: {e}")
    DOCX_AVAILABLE = False

def test_hyperlink_paragraph_state():
    """Test if paragraph state affects hyperlink creation."""
    
    if not DOCX_AVAILABLE:
        print("❌ Cannot test without python-docx")
        return
    
    print("🧪 Testing Hyperlink Paragraph State")
    print("=" * 50)
    
    # Create test document
    doc = Document()
    styles = DocumentStyles()
    
    # Test skills data (same as API)
    test_skills = """• Stakeholder Requirements|https://lightcast.io/open-skills/skills/ES21E5D7885B46FE1773
• Growth Planning|https://lightcast.io/open-skills/skills/ES42E5D7967FCC7EDD0D
• Vision Development|https://lightcast.io/open-skills/skills/ES465F92F433C4A44E49"""
    
    print(f"📋 Test skills data:")
    print(f"Raw: {repr(test_skills)}")
    
    # Create a test table
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    
    print(f"\n🔍 Testing Current _add_cell_content_with_hyperlinks Method")
    
    # Test current implementation
    styles._add_cell_content_with_hyperlinks(cell, test_skills)
    
    # Examine the paragraph content
    para = cell.paragraphs[0]
    print(f"📊 Paragraph has {len(para.runs)} runs:")
    
    for i, run in enumerate(para.runs):
        print(f"   Run {i+1}: {repr(run.text)}")
        print(f"      Font: {run.font.name}")
        print(f"      Underline: {run.font.underline}")
        print(f"      Color: {run.font.color.rgb if run.font.color.rgb else 'None'}")
        
        # Check if this is a hyperlink by examining the XML
        if hasattr(run._element, 'getparent'):
            parent = run._element.getparent()
            if parent is not None and parent.tag.endswith('hyperlink'):
                print(f"      ✅ IS HYPERLINK")
            else:
                print(f"      ❌ NOT HYPERLINK")
        else:
            print(f"      ❓ Cannot determine hyperlink status")
    
    # Save test document
    doc.save('hyperlink_test_current.docx')
    print(f"\n💾 Saved test document: hyperlink_test_current.docx")

def test_improved_hyperlink_method():
    """Test an improved version of the hyperlink method."""
    
    if not DOCX_AVAILABLE:
        print("❌ Cannot test without python-docx")
        return
    
    print(f"\n🔧 Testing Improved Hyperlink Method")
    print("=" * 50)
    
    # Create test document
    doc = Document()
    styles = DocumentStyles()
    
    # Test skills data (same as API)
    test_skills = """• Stakeholder Requirements|https://lightcast.io/open-skills/skills/ES21E5D7885B46FE1773
• Growth Planning|https://lightcast.io/open-skills/skills/ES42E5D7967FCC7EDD0D
• Vision Development|https://lightcast.io/open-skills/skills/ES465F92F433C4A44E49"""
    
    # Create a test table
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    
    # Test improved implementation
    _add_cell_content_with_hyperlinks_fixed(styles, cell, test_skills)
    
    # Examine the paragraph content
    para = cell.paragraphs[0]
    print(f"📊 Paragraph has {len(para.runs)} runs:")
    
    for i, run in enumerate(para.runs):
        print(f"   Run {i+1}: {repr(run.text)}")
        print(f"      Font: {run.font.name}")
        print(f"      Underline: {run.font.underline}")
        print(f"      Color: {run.font.color.rgb if run.font.color.rgb else 'None'}")
        
        # Check if this is a hyperlink by examining the XML
        if hasattr(run._element, 'getparent'):
            parent = run._element.getparent()
            if parent is not None and parent.tag.endswith('hyperlink'):
                print(f"      ✅ IS HYPERLINK")
            else:
                print(f"      ❌ NOT HYPERLINK")
        else:
            print(f"      ❓ Cannot determine hyperlink status")
    
    # Save test document
    doc.save('hyperlink_test_fixed.docx')
    print(f"\n💾 Saved test document: hyperlink_test_fixed.docx")

def _add_cell_content_with_hyperlinks_fixed(styles, cell, content: str):
    """FIXED version of _add_cell_content_with_hyperlinks method."""
    
    para = cell.paragraphs[0]
    para.clear()
    
    print(f"🔧 FIXED: Processing content: {repr(content[:100])}...")
    
    # Check if content contains skills with URLs
    if '•' in content and '|https://lightcast.io/' in content:
        # Split by newlines to handle multiple skills
        skills = content.split('\n')
        print(f"🔧 FIXED: Found {len(skills)} skills")
        
        for i, skill in enumerate(skills):
            skill = skill.strip()
            if not skill:
                print(f"🔧 FIXED: Skill {i+1}: EMPTY - skipping")
                continue
            
            print(f"🔧 FIXED: Processing skill {i+1}: {repr(skill)}")
            
            # CRITICAL FIX: Always add line break AFTER hyperlink, not before
            # This ensures the first skill gets processed without paragraph state issues
            
            # Check if this skill has an embedded URL
            if '|https://lightcast.io/' in skill:
                # Extract skill name and URL
                parts = skill.split('|', 1)
                skill_name = parts[0].strip()
                skill_url = parts[1].strip()
                
                print(f"🔧 FIXED: Creating hyperlink: '{skill_name}' -> '{skill_url}'")
                
                # Add hyperlink for this skill
                styles._add_hyperlink(para, skill_name, skill_url)
                
                # Add line break AFTER the hyperlink (except for last skill)
                if i < len(skills) - 1:
                    print(f"🔧 FIXED: Adding line break after skill {i+1}")
                    para.add_run('\n')
                
            else:
                # Regular skill without URL
                print(f"🔧 FIXED: Adding regular text: '{skill}'")
                run = para.add_run(skill)
                run.font.name = styles.fonts.FONT_PRIMARY
                run.font.size = styles.fonts.get_pt_size('table_body')
                if styles.colors.get_rgb_color('nab_black'):
                    run.font.color.rgb = styles.colors.get_rgb_color('nab_black')
                
                # Add line break AFTER the text (except for last skill)
                if i < len(skills) - 1:
                    para.add_run('\n')
    else:
        # Regular cell content without skills
        print(f"🔧 FIXED: Adding regular content: {repr(content[:50])}...")
        run = para.add_run(content)
        run.font.name = styles.fonts.FONT_PRIMARY
        run.font.size = styles.fonts.get_pt_size('table_body')
        if styles.colors.get_rgb_color('nab_black'):
            run.font.color.rgb = styles.colors.get_rgb_color('nab_black')

def test_hyperlink_xml_structure():
    """Test the actual XML structure of hyperlinks."""
    
    if not DOCX_AVAILABLE:
        print("❌ Cannot test without python-docx")
        return
    
    print(f"\n🔍 Testing Hyperlink XML Structure")
    print("=" * 50)
    
    # Create test document
    doc = Document()
    styles = DocumentStyles()
    
    # Create a simple paragraph to test hyperlink creation
    para = doc.add_paragraph()
    
    # Test the _add_hyperlink method directly
    print(f"🧪 Testing _add_hyperlink method directly")
    
    try:
        styles._add_hyperlink(para, "• Test Skill", "https://lightcast.io/open-skills/skills/TEST123")
        print(f"✅ Hyperlink creation succeeded")
        
        # Examine the XML
        print(f"📋 Paragraph XML:")
        print(para._element.xml)
        
    except Exception as e:
        print(f"❌ Hyperlink creation failed: {e}")

def main():
    """Run all hyperlink tests."""
    
    print("🚨 NAB Skills Intelligence Platform - Hyperlink Fix Test")
    print("🎯 Goal: Fix the first skill hyperlink issue")
    print("")
    
    if not DOCX_AVAILABLE:
        print("❌ python-docx not available, cannot run tests")
        return
    
    # Test current implementation
    test_hyperlink_paragraph_state()
    
    # Test improved implementation
    test_improved_hyperlink_method()
    
    # Test XML structure
    test_hyperlink_xml_structure()
    
    print(f"\n🎯 Compare the two generated documents:")
    print("1. hyperlink_test_current.docx - Current implementation")
    print("2. hyperlink_test_fixed.docx - Fixed implementation")
    print("")
    print("🔍 Check if the first skill in both documents has hyperlink formatting")

if __name__ == "__main__":
    main() 
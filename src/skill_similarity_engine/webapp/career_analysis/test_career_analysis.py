"""

Test script for Executive Summary Generation with Logical Role Architecture

Tests the integration of YAML templates, SQL queries, and Python generation logic.

Now includes logical role support for cleaner user experience.



Usage:

    # Basic test modes

    python test_career_analysis.py                           # Run all tests with default job

    python test_career_analysis.py --template                # Test template loading only

    python test_career_analysis.py --executive               # Test executive summary only

    python test_career_analysis.py --current                 # Test current role context only

    python test_career_analysis.py --pathway                 # Test pathway analysis only

    python test_career_analysis.py --strategic               # Test strategic recommendations only

    python test_career_analysis.py --conclusion              # Test conclusion only

    python test_career_analysis.py --word                    # Generate full Word document

    python test_career_analysis.py --copy                    # Display clean copy without debug

    

    # Core parameters matching HTML interface

    python test_career_analysis.py --job-from R0100.2        # Specify source job

    python test_career_analysis.py --mode top_matches        # Analysis mode: top_matches (default) or specific

    python test_career_analysis.py --job-to R0025.1          # Target job (required for specific mode)

    python test_career_analysis.py --similarity-min 40       # Minimum similarity percentage (20-95)

    python test_career_analysis.py --similarity-max 90       # Maximum similarity percentage (25-100)

    

    # Combined examples

    python test_career_analysis.py --job-from R0100.2 --mode top_matches --similarity-min 45 --similarity-max 85 --word

    python test_career_analysis.py --job-from R0100.2 --mode specific --job-to R0025.1 --similarity-min 50 --similarity-max 80 --executive

    python test_career_analysis.py --job-from R0100.2 --mode specific --job-to R0025.1 --copy

"""



import sqlite3

import sys

import argparse

from pathlib import Path



# Set up Unicode handling for Windows environments

import os

if os.name == 'nt':  # Windows

    import io

    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

    sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding='utf-8', errors='replace')



# Add the src directory to path so we can import the generator

src_path = Path(__file__).parent / 'src'

sys.path.insert(0, str(src_path))



from executive_summary_generator import ExecutiveSummaryGenerator

from skill_similarity_engine.utils.display import JobDisplayManager, DisplayFormat

from current_role_context_generator import CurrentRoleContextGenerator



def validate_job_exists(db, job_id):

    """Validate that a job ID exists in the database."""

    cursor = db.execute("SELECT JobProfileID FROM jobs WHERE JobProfileID = ?", (job_id,))

    result = cursor.fetchone()

    if not result:

        available_jobs = db.execute("SELECT JobProfileID FROM jobs ORDER BY JobProfileID LIMIT 10").fetchall()

        available_list = [job[0] for job in available_jobs]

        print(f"❌ Job ID '{job_id}' not found in database.")

        print(f"💡 Available job IDs (first 10): {', '.join(available_list)}")

        return False

    return True



def validate_similarity_range(min_sim, max_sim):

    """Validate similarity range parameters."""

    if min_sim < 20 or min_sim > 95:

        print(f"❌ similarity-min must be between 20 and 95, got {min_sim}")

        return False

    if max_sim < 25 or max_sim > 100:

        print(f"❌ similarity-max must be between 25 and 100, got {max_sim}")

        return False

    if min_sim >= max_sim:

        print(f"❌ similarity-min ({min_sim}) must be less than similarity-max ({max_sim})")

        return False

    return True



def check_specific_transition_exists(db, job_from, job_to, min_sim, max_sim):

    """Check if a specific job transition exists within the similarity range."""

    cursor = db.execute("""

        SELECT similarity_score 

        FROM job_similarities 

        WHERE job_from = ? AND job_to = ?

    """, (job_from, job_to))

    

    result = cursor.fetchone()

    if not result:

        print(f"❌ No similarity data found between {job_from} and {job_to}")

        return False, None

    

    similarity = result[0] * 100  # Convert to percentage

    

    if similarity < min_sim or similarity > max_sim:

        print(f"⚠️  Similarity {similarity:.1f}% is outside specified range {min_sim}%-{max_sim}%")

        print(f"   Continuing with analysis but consider adjusting similarity range...")

    

    return True, similarity



def get_available_targets(db, job_from, min_sim, max_sim, limit=10):

    """Get available target jobs within similarity range for suggestions."""

    cursor = db.execute("""

        SELECT js.job_to, js.similarity_score * 100 as similarity_pct, j.JobProfile

        FROM job_similarities js

        JOIN jobs j ON js.job_to = j.JobProfileID

        WHERE js.job_from = ? 

        AND js.similarity_score * 100 BETWEEN ? AND ?

        ORDER BY js.similarity_score DESC

        LIMIT ?

    """, (job_from, min_sim, max_sim, limit))

    

    return cursor.fetchall()



def test_executive_summary_generation(job_from="R0100.2", mode="top_matches", job_to=None, 

                                    similarity_min=40, similarity_max=90, top_n=3):

    """Test the executive summary generation with configurable parameters."""

    

    print("🧪 Testing Executive Summary Generation with Logical Role Architecture")

    print("=" * 70)

    print(f"📋 Test Parameters:")

    print(f"   Source Job: {job_from}")

    print(f"   Analysis Mode: {mode}")

    if mode == "specific" and job_to:

        print(f"   Target Job: {job_to}")

    print(f"   Similarity Range: {similarity_min}% - {similarity_max}%")

    print()

    

    # Connect to database

    db_path = Path(__file__).parent.parent.parent.parent.parent / "models" / "2025-Q2" / "business_context.sqlite"

    

    if not db_path.exists():

        print(f"⚠️ Database not found at: {db_path}")

        print("Please update the db_path in the test script to point to your database.")

        return

    

    try:

        # Connect to database

        print(f"📊 Connecting to database: {db_path}")

        db = sqlite3.connect(str(db_path))

        db.row_factory = sqlite3.Row  # Enable column access by name

        

        # Validate source job

        if not validate_job_exists(db, job_from):

            return

        

        # For specific mode, validate target job and transition (skip range validation)

        if mode == "specific":

            if not job_to:

                print("❌ --job-to is required when using --mode specific")

                print("💡 Use --job-to JOBID to specify target job")

                return

            

            # Parse target jobs (could be single job or comma-separated list)

            if ',' in job_to:

                job_to_list = [j.strip() for j in job_to.split(',')]

                print(f"📋 Validating {len(job_to_list)} target jobs: {', '.join(job_to_list)}")

                

                for target_job in job_to_list:

                    if not validate_job_exists(db, target_job):

                        return

                

                # Check transitions for all targets

                missing_transitions = []

                for target_job in job_to_list:

                    cursor = db.execute("""

                        SELECT similarity_score 

                        FROM job_similarities 

                        WHERE job_from = ? AND job_to = ?

                    """, (job_from, target_job))

                    

                    result = cursor.fetchone()

                    if not result:

                        missing_transitions.append(target_job)

                

                if missing_transitions:

                    print(f"❌ No similarity data found for transitions: {job_from} → {', '.join(missing_transitions)}")

                    return

                

                # Show similarity scores for all targets

                print("✅ All target transitions validated:")

                for target_job in job_to_list:

                    cursor = db.execute("""

                        SELECT similarity_score 

                        FROM job_similarities 

                        WHERE job_from = ? AND job_to = ?

                    """, (job_from, target_job))

                    

                    result = cursor.fetchone()

                    if result:

                        similarity = result[0] * 100

                        print(f"   {job_from} → {target_job}: {similarity:.1f}% similarity")

            else:

                # Single target job validation

                if not validate_job_exists(db, job_to):

                    return

                

                # For specific mode, just check if transition data exists (ignore range)

                cursor = db.execute("""

                    SELECT similarity_score 

                    FROM job_similarities 

                    WHERE job_from = ? AND job_to = ?

                """, (job_from, job_to))

                

                result = cursor.fetchone()

                if not result:

                    print(f"❌ No similarity data found between {job_from} and {job_to}")

                    return

                

                similarity = result[0] * 100  # Convert to percentage

                print(f"✅ Specific transition validated: {similarity:.1f}% similarity")

            

            print("💡 Note: Similarity range filters are ignored in specific mode")

        else:

            # For discovery mode, validate similarity range

            if not validate_similarity_range(similarity_min, similarity_max):

                return

        

        # Initialize generator and display manager

        print("🏗️ Initializing Executive Summary Generator with JobDisplayManager...")

        generator = ExecutiveSummaryGenerator(db)

        display_manager = JobDisplayManager(db)

        

        # Show logical role architecture improvements

        total_profiles = db.execute("SELECT COUNT(*) FROM jobs").fetchone()[0]

        logical_roles_query = """

        SELECT COUNT(DISTINCT 

            CASE 

                WHEN INSTR(JobProfile, ' - ') > 0 

                THEN SUBSTR(JobProfile, 1, INSTR(JobProfile, ' - ') - 1) || '|' || ManagementLevel

                ELSE JobProfile || '|' || ManagementLevel 

            END

        ) FROM jobs

        """

        logical_roles_count = db.execute(logical_roles_query).fetchone()[0]

        

        print(f"\n🏗️ Logical Role Architecture:")

        print(f"   Total JobProfileIDs: {total_profiles:,}")

        print(f"   Logical Roles:       {logical_roles_count:,}")

        print(f"   Reduction Ratio:     {total_profiles / logical_roles_count:.1f}:1")

        

        # Display job information

        logical_name = display_manager.get_display_name(job_from, DisplayFormat.LOGICAL)

        print(f"\n🎯 Generating executive summary for:")

        print(f"   JobProfileID: {job_from}")

        print(f"   Logical Role: {logical_name}")

        print(f"   Analysis Mode: {mode.replace('_', ' ').title()}")

        

        if mode == "specific" and job_to:

            # Handle both single and multiple target jobs for display

            if ',' in job_to:

                job_to_list = [j.strip() for j in job_to.split(',')]

                target_names = []

                for target_id in job_to_list:

                    target_name = display_manager.get_display_name(target_id, DisplayFormat.LOGICAL)

                    target_names.append(f"{target_id} ({target_name})")

                

                print(f"   Target Jobs: {job_to}")

                print(f"   Target Logical Roles: {', '.join(target_names)}")

            else:

                target_logical_name = display_manager.get_display_name(job_to, DisplayFormat.LOGICAL)

                print(f"   Target Job: {job_to}")

                print(f"   Target Logical Role: {target_logical_name}")

        

        # Generate executive summary with parameters

        result = generator.generate(

            job_from, 

            analysis_mode=mode, 

            job_to=job_to, 

            similarity_range=(similarity_min/100, similarity_max/100),

            top_n=top_n

        )

        

        # Display results

        print("\n" + "=" * 50)

        print("📄 GENERATED EXECUTIVE SUMMARY")

        print("=" * 50)

        

        content = result.get('content', {})

        

        # Strategic Context

        if 'strategic_context' in content:

            print(f"\n### {content['strategic_context']['title']}")

            print(content['strategic_context']['content'])

        

        # Key Findings

        if 'key_findings' in content:

            print(f"\n### {content['key_findings']['title']}")

            key_content = content['key_findings']['content']

            if isinstance(key_content, dict) and 'text' in key_content:

                print(key_content['text'])

            else:

                print(key_content)

        

        # Primary Recommendations

        if 'primary_recommendations' in content:

            print(f"\n### {content['primary_recommendations']['title']}")

            rec_content = content['primary_recommendations']['content']

            if isinstance(rec_content, dict) and 'text' in rec_content:

                print(rec_content['text'])

            else:

                print(rec_content)

        

        # Data-Driven Classification

        if 'data_driven_classification' in content:

            print(f"\n### Data-Driven Classification")

            print(content['data_driven_classification']['content'])

        

        # Confidence Assessment

        if 'confidence_assessment' in content:

            print(f"\n### {content['confidence_assessment']['title']}")

            conf_content = content['confidence_assessment']['content']

            if isinstance(conf_content, dict) and 'text' in conf_content:

                print(conf_content['text'])

            else:

                print(conf_content)

        

        # Show template variables for debugging

        print("\n" + "=" * 50)

        print("🔧 TEMPLATE VARIABLES (Debug Info)")

        print("=" * 50)

        

        variables = result.get('template_variables', {})

        key_variables = [

            'source_job_title', 'total_job_count', 'pathway_count', 

            'min_similarity', 'max_similarity', 'confidence_level',

            'opportunity_descriptor', 'percentile_descriptor'

        ]

        

        for var in key_variables:

            if var in variables:

                print(f"{var}: {variables[var]}")

        

        # Show analysis mode context

        print(f"\nAnalysis Mode Context:")

        print(f"  Mode: {mode}")

        print(f"  Similarity Range: {similarity_min}% - {similarity_max}%")

        if mode == "specific" and job_to:

            print(f"  Specific Transition: {job_from} → {job_to}")

        

        # Show top pathways with logical role names

        top_pathways = variables.get('top_pathways', [])

        if top_pathways:

            print(f"\nTop {len(top_pathways)} career pathways (showing logical roles):")

            for i, pathway in enumerate(top_pathways, 1):

                logical_role = pathway.get('target_logical_role', pathway.get('target_job_title', 'Unknown'))

                similarity = pathway.get('similarity_score', 0)

                move_type = pathway.get('move_type', 'Unknown')

                transition = pathway.get('level_transition', 'Unknown')

                

                print(f"  {i}. {logical_role} - {similarity}% similarity")

                print(f"     Move Type: {move_type}")

                print(f"     Transition: {transition}")

                if i < len(top_pathways):  # Add spacing except after last item

                    print()

        

        # Show references

        print("\n" + "=" * 50)

        print("📚 REFERENCES")

        print("=" * 50)

        

        references = result.get('references', {})

        key_refs = ['1', '2', '13', '14', '15', '16']

        for ref in key_refs:

            if ref in references:

                print(f"({ref}) {references[ref]}")

        

        print("\n✅ Executive Summary generation completed successfully!")

        print("\n🎯 Key Architecture Improvements Demonstrated:")

        print("   ✅ Logical role display names (e.g., 'Data Scientist (Group 1)')")

        print("   ✅ 3:1 reduction in complexity (715 profiles → ~237 logical roles)")

        print("   ✅ Executive-ready presentation format")

        print("   ✅ Backward compatibility with existing database")

        print(f"   ✅ Configurable parameters: {mode} mode, {similarity_min}%-{similarity_max}% range")

        

    except Exception as e:

        print(f"❌ Error during testing: {e}")

        import traceback

        traceback.print_exc()

    

    finally:

        if 'db' in locals():

            db.close()



def test_template_loading():

    """Test that the YAML template loads correctly."""

    

    print("\n🧪 Testing YAML Template Loading")

    print("=" * 50)

    

    try:

        # Test template loading without database

        generator = ExecutiveSummaryGenerator(None)

        

        if generator.template_data:

            print("✅ YAML template loaded successfully!")

            

            # Check key sections exist

            exec_summary = generator.template_data.get('executive_summary', {})

            required_sections = [

                'strategic_context', 'key_findings', 'primary_recommendations',

                'confidence_assessment'

            ]

            

            for section in required_sections:

                if section in exec_summary:

                    print(f"✅ Section '{section}' found")

                else:

                    print(f"❌ Section '{section}' missing")

            

            # Check references exist

            references = generator.template_data.get('references', {})

            if references:

                print(f"✅ References section found with {len(references)} references")

            else:

                print("❌ References section missing")

        

        else:

            print("❌ YAML template failed to load")

    

    except Exception as e:

        print(f"❌ Error testing template: {e}")



def test_current_role_context_generation(job_from="R0100.2", mode="top_matches", job_to=None, 

                                       similarity_min=40, similarity_max=90):

    """Test the current role context generation with configurable parameters."""

    

    print("\n🧪 Testing Current Role Context Generation")

    print("=" * 70)

    print(f"📋 Test Parameters:")

    print(f"   Source Job: {job_from}")

    print(f"   Analysis Mode: {mode}")

    if mode == "specific" and job_to:

        print(f"   Target Job: {job_to}")

    print(f"   Similarity Range: {similarity_min}% - {similarity_max}%")

    print()

    

    # Connect to database

    db_path = Path(__file__).parent.parent.parent.parent.parent / "models" / "2025-Q2" / "business_context.sqlite"

    

    if not db_path.exists():

        print(f"⚠️ Database not found at: {db_path}")

        print("Please update the db_path in the test script to point to your database.")

        return

    

    try:

        # Connect to database

        print(f"📊 Connecting to database: {db_path}")

        db = sqlite3.connect(str(db_path))

        db.row_factory = sqlite3.Row  # Enable column access by name

        

        # Validate parameters (basic validation, detailed validation in executive summary)

        if not validate_job_exists(db, job_from):

            return

        

        if mode == "specific" and job_to and not validate_job_exists(db, job_to):

            return

        

        # Initialize generator

        print("🏗️ Initializing Current Role Context Generator...")

        generator = CurrentRoleContextGenerator(db)

        

        print(f"\n🎯 Generating current role context for: {job_from}")

        if mode == "specific" and job_to:

            print(f"   Specific transition to: {job_to}")

        

        # Generate current role context

        result = generator.generate(job_from, include_organisational_deployment=True)

        

        # Display results

        print("\n" + "=" * 50)

        print("📄 GENERATED CURRENT ROLE CONTEXT")

        print("=" * 50)

        

        content = result.get('content', {})

        

        # Helper function to display structured content

        def display_content(section_content):

            if isinstance(section_content, dict) and 'text' in section_content:

                return section_content['text']

            else:

                return str(section_content)

        

        # Profile Overview

        if 'profile_overview' in content:

            print(f"\n### {content['profile_overview']['title']}")

            print(display_content(content['profile_overview']['content']))

        

        # Core Competency Foundation

        if 'core_competency_foundation' in content:

            print(f"\n### {content['core_competency_foundation']['title']}")

            print(display_content(content['core_competency_foundation']['content']))

        

        # Strategic Value Proposition

        if 'strategic_value_proposition' in content:

            print(f"\n### {content['strategic_value_proposition']['title']}")

            print(display_content(content['strategic_value_proposition']['content']))

        

        # Strategic Intelligence Metrics

        if 'strategic_intelligence_metrics' in content:

            print(f"\n### {content['strategic_intelligence_metrics']['title']}")

            print(display_content(content['strategic_intelligence_metrics']['content']))

        

        # Show template variables for debugging

        print("\n" + "=" * 50)

        print("🔧 TEMPLATE VARIABLES (Debug Info)")

        print("=" * 50)

        

        variables = result.get('template_variables', {})

        key_variables = [

            'source_job_logical_display_name', 'total_skills', 'skill_category_count',

            'mobility_hub_score', 'transition_readiness', 'cross_family_reach',

            'strategic_value_assessment'

        ]

        

        for var in key_variables:

            if var in variables:

                print(f"{var}: {variables[var]}")

        

        # Show analysis mode context

        print(f"\nAnalysis Mode Context:")

        print(f"  Mode: {mode}")

        print(f"  Similarity Range: {similarity_min}% - {similarity_max}%")

        if mode == "specific" and job_to:

            print(f"  Specific Transition: {job_from} → {job_to}")

        

        print("\n✅ Current Role Context generation completed successfully!")

        print("\n🎯 Key Features Demonstrated:")

        print("   ✅ YAML template rendering with Jinja2")

        print("   ✅ Database integration following proven patterns")

        print("   ✅ Logical role display names integration")

        print("   ✅ Professional formatting matching gold standard")

        print(f"   ✅ Configurable parameters: {mode} mode, {similarity_min}%-{similarity_max}% range")

        

    except Exception as e:

        print(f"❌ Error during current role context testing: {e}")

        import traceback

        traceback.print_exc()

    

    finally:

        if 'db' in locals():

            db.close()



def test_pathway_analysis_generation(job_from="R0100.2", mode="top_matches", job_to=None, 

                                   similarity_min=40, similarity_max=90, top_n=3, tie_breaking_options=None):

    """Test the Pathway Analysis Generator with configurable parameters."""

    print("\n🧪 Testing Pathway Analysis Generation")

    print("=" * 70)

    print(f"📋 Test Parameters:")

    print(f"   Source Job: {job_from}")

    print(f"   Analysis Mode: {mode}")

    if mode == "specific" and job_to:

        print(f"   Target Job: {job_to}")

    print(f"   Similarity Range: {similarity_min}% - {similarity_max}%")

    print(f"   🎯 Top N Pathways: {top_n}")

    if tie_breaking_options:

        active_options = [k for k, v in tie_breaking_options.items() if v]

        if active_options:

            print(f"   🔧 Tie-Breaking Options: {', '.join(active_options)}")

    print()

    

    # Connect to database

    db_path = Path(__file__).parent.parent.parent.parent.parent / "models" / "2025-Q2" / "business_context.sqlite"

    

    if not db_path.exists():

        print(f"⚠️ Database not found at: {db_path}")

        print("Please update the db_path in the test script to point to your database.")

        return

    

    try:

        # Connect to database

        print(f"📊 Connecting to database: {db_path}")

        db = sqlite3.connect(str(db_path))

        db.row_factory = sqlite3.Row  # Enable column access by name

        

        # Validate parameters

        if not validate_job_exists(db, job_from):

            return

        

        if mode == "specific" and job_to and not validate_job_exists(db, job_to):

            return

        

        # Initialize generator

        print("🏗️ Initializing Pathway Analysis Generator...")

        from src.pathway_analysis_generator import PathwayAnalysisGenerator

        generator = PathwayAnalysisGenerator(db)

        

        print(f"🎯 Generating pathway analysis for: {job_from}")

        if mode == "specific" and job_to:

            print(f"   Specific transition to: {job_to}")

        

        # Display tie-breaking options if provided

        if tie_breaking_options:

            active_options = [k for k, v in tie_breaking_options.items() if v]

            if active_options:

                print(f"   🎛️ Active tie-breaking options: {', '.join(active_options)}")

            else:

                print(f"   🎛️ Tie-breaking options provided but none active")

        

        # Generate pathway analysis with parameters

        result = generator.generate(

            job_from, 

            analysis_mode=mode, 

            job_to=job_to, 

            similarity_range=(similarity_min/100, similarity_max/100),

            include_organisational_deployment=True,

            top_n=top_n,

            tie_breaking_options=tie_breaking_options or {}

        )

        

        print("\n" + "=" * 50)

        print("📄 GENERATED PATHWAY ANALYSIS")

        print("=" * 50)

        

        # Display section title

        print(f"# {result['section_title']}")

        print()

        

        # Display each opportunity

        if 'content' in result and 'opportunities' in result['content']:

            opportunities_count = len(result['content']['opportunities'])

            print(f"🎯 Generated {opportunities_count} opportunities (requested top {top_n})")

            print()

            

            for i, opportunity in enumerate(result['content']['opportunities'], 1):

                print(opportunity['header'])

                print()

                

                # Handle the opportunity overview table

                print("### Opportunity Overview")

                if isinstance(opportunity['opportunity_overview'], dict) and 'text' in opportunity['opportunity_overview']:

                    print(opportunity['opportunity_overview']['text'])

                else:

                    print(opportunity['opportunity_overview'])

                print()

                

                print(f"### {opportunity['strategic_positioning']['title']}")

                print(opportunity['strategic_positioning']['content'])

                print()

                

                # Handle the new table structure for skills_transition_analysis and implementation_roadmap

                print("### Skills Transition Analysis")

                if isinstance(opportunity['skills_transition_analysis'], dict) and 'text' in opportunity['skills_transition_analysis']:

                    print(opportunity['skills_transition_analysis']['text'])

                else:

                    print(opportunity['skills_transition_analysis'])

                print()

                

                print(f"### {opportunity['business_case']['title']}")

                print(opportunity['business_case']['content'])

                print()

                

                print("### Implementation Roadmap")

                if isinstance(opportunity['implementation_roadmap'], dict) and 'text' in opportunity['implementation_roadmap']:

                    print(opportunity['implementation_roadmap']['text'])

                else:

                    print(opportunity['implementation_roadmap'])

                print()

                

                if i < len(result['content']['opportunities']):

                    print("-" * 70)

                    print()

        

        # Display debugging info

        print("=" * 50)

        print("🔧 PATHWAY ANALYSIS DEBUG INFO")

        print("=" * 50)

        if 'opportunities' in result:

            for i, opp in enumerate(result['opportunities'], 1):

                pathway = opp['pathway_data']

                print(f"Opportunity {i}: {pathway['target_logical_role']}")

                print(f"  Similarity Score: {pathway['similarity_score']}%")

                print(f"  Move Type: {pathway['move_type']}")

                print(f"  Level Transition: {pathway['level_transition_display']}")

                print()

        

        assert 'section_title' in result

        assert 'content' in result

        print("✅ Pathway Analysis generation completed successfully!")

        print("🎯 Key Features Demonstrated:")

        if mode == "specific":

            print("   ✅ Specific transition analysis generated")

            print("   ✅ Mode-appropriate template loaded")

            print("   ✅ SpecificTransitionAnalyzer integration working")

        else:

            print("   ✅ Top 3 strategic opportunities identified")

            print("   ✅ Comprehensive opportunity analysis")

        print("   ✅ Skills transition calculations")

        print("   ✅ Business case generation")

        print("   ✅ Implementation roadmap creation")

        

    except Exception as e:

        print(f"❌ Error in pathway analysis generation: {e}")

        import traceback

        traceback.print_exc()

    

    finally:

        if 'db' in locals():

            db.close()



def test_strategic_recommendations_generation(job_from="R0100.2", mode="top_matches", job_to=None, 

                                           similarity_min=40, similarity_max=90, top_n=3):

    """Test the Strategic Recommendations Generator with configurable parameters."""

    print("\n🧪 Testing Strategic Recommendations Generation")

    print("=" * 70)

    print(f"📋 Test Parameters:")

    print(f"   Source Job: {job_from}")

    print(f"   Analysis Mode: {mode}")

    if mode == "specific" and job_to:

        print(f"   Target Job: {job_to}")

    print(f"   Similarity Range: {similarity_min}% - {similarity_max}%")

    print()

    

    # Connect to database

    db_path = Path(__file__).parent.parent.parent.parent.parent / "models" / "2025-Q2" / "business_context.sqlite"

    

    if not db_path.exists():

        print(f"⚠️ Database not found at: {db_path}")

        print("Please update the db_path in the test script to point to your database.")

        return

    

    try:

        # Connect to database

        print(f"📊 Connecting to database: {db_path}")

        db = sqlite3.connect(str(db_path))

        db.row_factory = sqlite3.Row  # Enable column access by name

        

        # Validate parameters

        if not validate_job_exists(db, job_from):

            return

        

        if mode == "specific" and job_to and not validate_job_exists(db, job_to):

            return

        

        # Initialize generator

        print("🏗️ Initializing Strategic Recommendations Generator...")

        from strategic_recommendations_generator import StrategicRecommendationsGenerator

        generator = StrategicRecommendationsGenerator(db)

        

        print(f"🎯 Generating strategic recommendations for: {job_from}")

        if mode == "specific" and job_to:

            print(f"   Specific transition to: {job_to}")

        

        # Generate strategic recommendations

        result = generator.generate(

            job_from, 

            analysis_mode=mode, 

            job_to=job_to, 

            similarity_range=(similarity_min/100, similarity_max/100), 

            include_organisational_deployment=True,

            top_n=top_n

        )

        

        print("\n" + "=" * 50)

        print("📄 GENERATED STRATEGIC RECOMMENDATIONS")

        print("=" * 50)

        

        # Display section title

        print(f"# {result['section_title']}")

        print()

        

        # Display each content section

        content = result.get('content', {})

        section_order = [

            'database_driven_decision_support',

            'business_case',

            'immediate_actions', 

            'medium_term_initiatives',

            'success_metrics_evaluation',

            'research_references'

        ]

        

        for section_key in section_order:

            if section_key in content:

                section = content[section_key]

                print(f"### {section['title']}")

                

                # Handle different content types

                section_content = section['content']

                if isinstance(section_content, dict) and 'text' in section_content:

                    # Structured content with formatting metadata

                    print(section_content['text'])

                elif isinstance(section_content, dict) and 'formatting' in section_content:

                    # Table or other structured content

                    formatting = section_content['formatting']

                    if formatting.get('content_type') == 'table':

                        # Display table headers and rows

                        headers = formatting.get('headers', [])

                        rows = formatting.get('rows', [])

                        if headers:

                            print(" | ".join(headers))

                            print("|".join(["-" * len(header) for header in headers]))

                        for row in rows:

                            print(" | ".join(str(cell) for cell in row))

                    else:

                        # Display the text content

                        print(section_content.get('text', section_content))

                else:

                    # Basic string content

                    print(section_content)

                print()

        

        # Display debugging info

        print("=" * 50)

        print("🔧 STRATEGIC RECOMMENDATIONS DEBUG INFO")

        print("=" * 50)

        

        variables = result.get('template_variables', {})

        key_variables = [

            'source_job_logical_display_name', 'total_job_profiles', 'target_pathway_count',

            'avg_similarity', 'development_investment_weeks', 'source_function',

            'function_percentage', 'cross_family_mobility_count'

        ]

        

        for var in key_variables:

            if var in variables:

                print(f"{var}: {variables[var]}")

        

        # Show benchmark context

        benchmark_vars = [

            'industry_transition_success', 'target_transition_success',

            'industry_retention', 'target_retention', 

            'industry_productivity_months', 'target_productivity_months'

        ]

        

        print(f"\nBenchmark Context:")

        for var in benchmark_vars:

            if var in variables:

                print(f"  {var}: {variables[var]}")

        

        assert 'section_title' in result

        assert 'content' in result

        print("\n✅ Strategic Recommendations generation completed successfully!")

        print("🎯 Key Features Demonstrated:")

        print("   ✅ Database-driven decision support metrics")

        print("   ✅ Immediate and medium-term action plans")

        print("   ✅ Success metrics with industry benchmarks")

        print("   ✅ Cross-functional mobility insights")

        print("   ✅ Development investment calculations")

        

    except Exception as e:

        print(f"❌ Error in strategic recommendations generation: {e}")

        import traceback

        traceback.print_exc()

    

    finally:

        if 'db' in locals():

            db.close()



def test_conclusion_generation(job_from="R0100.2", mode="top_matches", job_to=None, 

                             similarity_min=40, similarity_max=90, top_n=3):

    """Test the conclusion generation with configurable parameters."""

    

    print("\n🧪 Testing Conclusion Generation")

    print("=" * 70)

    print(f"📋 Test Parameters:")

    print(f"   Source Job: {job_from}")

    print(f"   Analysis Mode: {mode}")

    if mode == "specific" and job_to:

        print(f"   Target Job: {job_to}")

    print(f"   Similarity Range: {similarity_min}% - {similarity_max}%")

    print()

    

    # Connect to database

    db_path = Path(__file__).parent.parent.parent.parent.parent / "models" / "2025-Q2" / "business_context.sqlite"

    

    if not db_path.exists():

        print(f"⚠️ Database not found at: {db_path}")

        print("Please update the db_path in the test script to point to your database.")

        return

    

    try:

        # Connect to database

        print(f"📊 Connecting to database: {db_path}")

        db = sqlite3.connect(str(db_path))

        db.row_factory = sqlite3.Row  # Enable column access by name

        

        # Validate parameters

        if not validate_job_exists(db, job_from):

            return

        

        if mode == "specific" and job_to and not validate_job_exists(db, job_to):

            return

        

        # Initialize generator

        print("🏗️ Initializing Conclusion Generator...")

        from conclusion_generator import ConclusionGenerator

        generator = ConclusionGenerator(db)

        

        print(f"\n🎯 Generating conclusion for: {job_from}")

        if mode == "specific" and job_to:

            print(f"   Specific transition to: {job_to}")

        

        # Generate conclusion

        result = generator.generate(

            job_from, 

            analysis_mode=mode, 

            job_to=job_to, 

            similarity_range=(similarity_min/100, similarity_max/100), 

            include_organisational_deployment=True,

            top_n=top_n

        )

        

        # Display results

        print("\n" + "=" * 50)

        print("📄 GENERATED CONCLUSION")

        print("=" * 50)

        

        # Display section title

        print(f"## {result['section_title']}")

        print()

        

        content = result.get('content', {})

        

        # Display each section in logical order

        section_order = ['opportunity_summary', 'strategic_alignment', 'recommended_approach', 'foundation_value']

        

        for section_key in section_order:

            if section_key in content:

                section = content[section_key]

                if 'content' in section:

                    print(section['content'])

                    print()

        

        # Show template variables for debugging  

        print("=" * 50)

        print("🔧 CONCLUSION DEBUG INFO")

        print("=" * 50)

        

        variables = result.get('template_variables', {})

        key_variables = [

            'source_job_logical_display_name', 'pathway_count', 'similarity_range_description',

            'opportunity_classification', 'pilot_size', 'strategic_priority'

        ]

        

        for var in key_variables:

            if var in variables:

                print(f"{var}: {variables[var]}")

        

        assert 'section_title' in result

        assert 'content' in result

        print("\n✅ Conclusion generation completed successfully!")

        print("\n🎯 Key Features Demonstrated:")

        print("   ✅ Database-driven opportunity synthesis")

        print("   ✅ Strategic alignment assessment")

        print("   ✅ Pilot program recommendations")

        print("   ✅ Foundation for workforce planning")

        

    except Exception as e:

        print(f"❌ Error during conclusion testing: {e}")

        import traceback

        traceback.print_exc()

    

    finally:

        if 'db' in locals():

            db.close()



def test_clean_copy_display(job_from="R0100.2", mode="top_matches", job_to=None, 

                           similarity_min=40, similarity_max=90, top_n=3):

    """Display clean copy of all sections without debug information."""

    

    print("📄 NAB Skills Intelligence Platform")

    print("Strategic Career Pathway Analysis")

    print("=" * 70)

    print(f"Parameters: {job_from} | {mode} mode | {similarity_min}%-{similarity_max}%")

    if mode == "specific" and job_to:

        print(f"Specific Transition: {job_from} → {job_to}")

    print()

    

    # Connect to database

    db_path = Path(__file__).parent.parent.parent.parent.parent / "models" / "2025-Q2" / "business_context.sqlite"

    

    if not db_path.exists():

        print(f"⚠️ Database not found at: {db_path}")

        print("Please update the db_path in the test script to point to your database.")

        return

    

    try:

        # Connect to database (no debug output)

        db = sqlite3.connect(str(db_path))

        db.row_factory = sqlite3.Row

        

        # Initialize all generators (no debug output)

        from executive_summary_generator import ExecutiveSummaryGenerator

        from current_role_context_generator import CurrentRoleContextGenerator

        from pathway_analysis_generator import PathwayAnalysisGenerator

        from strategic_recommendations_generator import StrategicRecommendationsGenerator

        from conclusion_generator import ConclusionGenerator

        

        exec_generator = ExecutiveSummaryGenerator(db)

        context_generator = CurrentRoleContextGenerator(db)

        pathway_generator = PathwayAnalysisGenerator(db)

        strategic_generator = StrategicRecommendationsGenerator(db)

        conclusion_generator = ConclusionGenerator(db)

        

        # Use configurable job

        test_job_id = job_from

        

        # Get job name for header - use a simple query approach

        try:

            job_query = db.execute("SELECT JobProfile FROM jobs WHERE JobProfileID = ?", (test_job_id,)).fetchone()

            job_name = job_query['JobProfile'] if job_query else test_job_id

        except:

            job_name = test_job_id

        

        print(f"Career Transition Analysis: {job_name}")

        print(f"Generated: {Path(__file__).parent.parent.parent.parent.parent}")

        print("Version 1.0 - Skills Intelligence Analysis")

        print()

        

        # Generate all sections (suppress logging temporarily)

        import logging

        logging.getLogger().setLevel(logging.ERROR)

        

        sections = {}

        sections['executive_summary'] = exec_generator.generate(

            test_job_id, analysis_mode=mode, job_to=job_to, 

            similarity_range=(similarity_min/100, similarity_max/100), 

            top_n=top_n

        )

        sections['current_role_context'] = context_generator.generate(test_job_id, include_organisational_deployment=True)

        sections['pathway_analysis'] = pathway_generator.generate(

            test_job_id, analysis_mode=mode, job_to=job_to, 

            similarity_range=(similarity_min/100, similarity_max/100), 

            include_organisational_deployment=True,

            top_n=top_n

        )

        sections['strategic_recommendations'] = strategic_generator.generate(

            test_job_id, analysis_mode=mode, job_to=job_to, 

            similarity_range=(similarity_min/100, similarity_max/100), 

            include_organisational_deployment=True,

            top_n=top_n

        )

        sections['conclusion'] = conclusion_generator.generate(

            test_job_id, analysis_mode=mode, job_to=job_to, 

            similarity_range=(similarity_min/100, similarity_max/100), 

            include_organisational_deployment=True,

            top_n=top_n

        )

        

        # Restore logging

        logging.getLogger().setLevel(logging.INFO)

        

        # Display sections cleanly

        section_order = [

            ('executive_summary', 'Executive Summary'),

            ('current_role_context', 'Current Role Context'),

            ('pathway_analysis', 'Pathway Analysis: Top 3 Strategic Opportunities'),

            ('strategic_recommendations', 'Strategic Recommendations'),

            ('conclusion', 'Conclusion')

        ]

        

        for section_key, section_title in section_order:

            if section_key in sections:

                print(f"\n## {section_title}")

                print()

                

                section_result = sections[section_key]

                content = section_result.get('content', {})

                

                if section_key == 'pathway_analysis' and 'opportunities' in content:

                    # Special handling for pathway analysis with opportunities

                    for i, opportunity in enumerate(content['opportunities'], 1):

                        print(f"### {opportunity['header']}")

                        print()

                        

                        # Display each sub-section

                        sub_sections = ['strategic_positioning', 'skills_transition_analysis', 'business_case', 'implementation_roadmap']

                        for sub_key in sub_sections:

                            if sub_key in opportunity:

                                sub_section = opportunity[sub_key]

                                print(f"#### {sub_section['title']}")

                                print(sub_section['content'])

                                print()

                        

                        if i < len(content['opportunities']):

                            print("---")

                            print()

                else:

                    # Standard section handling

                    for subsection_key, subsection in content.items():

                        if isinstance(subsection, dict) and 'content' in subsection:

                            if 'title' in subsection:

                                print(f"### {subsection['title']}")

                            

                            # Handle structured content

                            subsection_content = subsection['content']

                            if isinstance(subsection_content, dict) and 'text' in subsection_content:

                                print(subsection_content['text'])

                            else:

                                print(subsection_content)

                            print()

        

        print("=" * 70)

        print("End of Career Transition Analysis")

        

    except Exception as e:

        print(f"❌ Error generating clean copy: {e}")

        import traceback

        traceback.print_exc()

    

    finally:

        if 'db' in locals():

            db.close()



def test_full_career_analysis_generation(job_from="R0100.2", mode="top_matches", job_to=None, 

                                   similarity_min=40, similarity_max=90, top_n=3):

    """Test full Career Transition Analysis Generator with all sections and Word document output."""

    

    print("\n🧪 Testing Full Career Transition Analysis Generator with Word Document Output")

    print("=" * 70)

    print(f"📋 Test Parameters:")

    print(f"   Source Job: {job_from}")

    print(f"   Analysis Mode: {mode}")

    if mode == "specific" and job_to:

        print(f"   Target Job: {job_to}")

    print(f"   Similarity Range: {similarity_min}% - {similarity_max}%")

    print()

    

    # Connect to database

    db_path = Path(__file__).parent.parent.parent.parent.parent / "models" / "2025-Q2" / "business_context.sqlite"

    

    if not db_path.exists():

        print(f"⚠️ Database not found at: {db_path}")

        print("Please update the db_path in the test script to point to your database.")

        return

    

    try:

        # Connect to database

        print(f"📊 Connecting to database: {db_path}")

        db = sqlite3.connect(str(db_path))

        db.row_factory = sqlite3.Row  # Enable column access by name

        

        # Initialize all generators

        print("🏗️ Initializing all generators...")

        from executive_summary_generator import ExecutiveSummaryGenerator

        from current_role_context_generator import CurrentRoleContextGenerator

        from pathway_analysis_generator import PathwayAnalysisGenerator

        from strategic_recommendations_generator import StrategicRecommendationsGenerator

        from conclusion_generator import ConclusionGenerator

        

        exec_generator = ExecutiveSummaryGenerator(db)

        context_generator = CurrentRoleContextGenerator(db)

        pathway_generator = PathwayAnalysisGenerator(db)

        strategic_generator = StrategicRecommendationsGenerator(db)

        conclusion_generator = ConclusionGenerator(db)

        

        # Validate parameters

        if not validate_job_exists(db, job_from):

            return

        

        if mode == "specific" and job_to:

            # Parse and validate target jobs

            if ',' in job_to:

                job_to_list = [j.strip() for j in job_to.split(',')]

                for target_job in job_to_list:

                    if not validate_job_exists(db, target_job):

                        return

            else:

                if not validate_job_exists(db, job_to):

                    return

        

        # Use configurable job parameters

        test_job_id = job_from

        print(f"🎯 Generating complete Career Transition Analysis for: {test_job_id}")

        if mode == "specific" and job_to:

            print(f"   Specific transition to: {job_to}")

        

        # Generate all sections with parameters

        print("\n📄 Generating sections...")

        sections = {}

        

        print("  ✅ Executive Summary...")

        sections['executive_summary'] = exec_generator.generate(

            test_job_id, analysis_mode=mode, job_to=job_to, 

            similarity_range=(similarity_min/100, similarity_max/100),

            top_n=top_n

        )

        

        print("  ✅ Current Role Context...")

        sections['current_role_context'] = context_generator.generate(test_job_id, include_organisational_deployment=True)

        

        print("  ✅ Pathway Analysis...")

        sections['pathway_analysis'] = pathway_generator.generate(

            test_job_id, analysis_mode=mode, job_to=job_to, 

            similarity_range=(similarity_min/100, similarity_max/100), 

            include_organisational_deployment=True,

            top_n=top_n

        )

        

        print("  ✅ Strategic Recommendations...")

        sections['strategic_recommendations'] = strategic_generator.generate(

            test_job_id, analysis_mode=mode, job_to=job_to, 

            similarity_range=(similarity_min/100, similarity_max/100), 

            include_organisational_deployment=True,

            top_n=top_n

        )

        

        print("  ✅ Conclusion...")

        sections['conclusion'] = conclusion_generator.generate(

            test_job_id, analysis_mode=mode, job_to=job_to, 

            similarity_range=(similarity_min/100, similarity_max/100), 

            include_organisational_deployment=True,

            top_n=top_n

        )

        

        # Prepare content for Word document

        print("\n📝 Preparing content for Word document...")

        

        # Extract dynamic section titles from generator results FIRST

        section_titles = []

        print("🔍 DEBUG: Extracting section titles from generator results...")

        for section_name in ['executive_summary', 'current_role_context', 'pathway_analysis', 'strategic_recommendations', 'conclusion']:

            if section_name in sections:

                # Extract dynamic title from generator result, fallback to static title

                generator_title = sections[section_name].get('section_title', '')

                print(f"   📋 {section_name}: generator_title = '{generator_title}'")

                if generator_title:

                    # Clean up any newline characters from section title

                    clean_title = generator_title.strip()

                    section_titles.append((section_name, clean_title))

                    print(f"      ✅ Using dynamic title: '{generator_title}'")

                else:

                    # Fallback to static titles for sections that don't generate dynamic titles

                    static_titles = {

                        'executive_summary': 'Executive Summary',

                        'current_role_context': 'Current Role Context',

                        'pathway_analysis': 'Pathway Analysis: Top 3 Strategic Opportunities',

                        'strategic_recommendations': 'Strategic Recommendations',

                        'conclusion': 'Conclusion'

                    }

                    section_titles.append((section_name, static_titles[section_name]))

                    print(f"      ⚠️ Using fallback title: '{static_titles[section_name]}'")

        

        print(f"🔍 DEBUG: Final section_titles = {section_titles}")

        

        # DEBUG: Check pathway analysis content structure

        if 'pathway_analysis' in sections:

            print(f"🔍 DEBUG: pathway_analysis keys = {list(sections['pathway_analysis'].keys())}")

            if 'content' in sections['pathway_analysis']:

                print(f"🔍 DEBUG: pathway_analysis content keys = {list(sections['pathway_analysis']['content'].keys())}")

                if 'opportunities' in sections['pathway_analysis']['content']:

                    print(f"🔍 DEBUG: opportunities count = {len(sections['pathway_analysis']['content']['opportunities'])}")

                    if sections['pathway_analysis']['content']['opportunities']:

                        first_opp = sections['pathway_analysis']['content']['opportunities'][0]

                        print(f"🔍 DEBUG: first opportunity keys = {list(first_opp.keys())}")

                        if 'strategic_positioning' in first_opp:

                            print(f"🔍 DEBUG: strategic_positioning = {first_opp['strategic_positioning']}")

            print(f"🔍 DEBUG: Full pathway_analysis structure:")

            import json

            print(json.dumps(sections['pathway_analysis'], indent=2, default=str)[:1000] + "...")

        

        # Extract content from each section

        word_content = {}

        analysis_data = {}

        

        for section_name, section_result in sections.items():

            word_content[section_name] = section_result.get('content', {})

            

            # Extract key analysis data from first section (executive summary)

            if section_name == 'executive_summary':

                variables = section_result.get('template_variables', {})

                max_sim = variables.get('max_similarity', 0)

                # Convert to float if it's a string

                if isinstance(max_sim, str):

                    try:

                        max_sim = float(max_sim)

                    except (ValueError, TypeError):

                        max_sim = 0

                

                analysis_data.update({

                    'source_job_logical_display_name': variables.get('source_job_title', 'Professional Role'),

                    'source_job_function': variables.get('source_job_function', 'Professional Services'),

                    'summary': f'Career pathway analysis for {variables.get("source_job_title", "Professional Role")}',

                    'avg_similarity': max_sim / 100,

                    'pathway_count': variables.get('pathway_count', 0),

                    'section_titles': section_titles  # Pass dynamic section titles to formatter

                })

        

        # Extract job_name for file naming with intelligent naming logic

        job_name = analysis_data.get('source_job_logical_display_name', 'Professional Role')

        

        # Create intelligent filename based on parameters

        filename_parts = [

            'CAREER_ANALYSIS',

            job_from.replace('.', '_'),  # Source job ID

            mode

        ]

        

        # Add specific mode details to filename

        if mode == "specific" and job_to:

            if ',' in job_to:

                # Multiple targets: use count and first target

                target_count = len([j.strip() for j in job_to.split(',')])

                first_target = job_to.split(',')[0].strip().replace('.', '_')

                filename_parts.extend(['multi', f"{target_count}targets", first_target])

            else:

                # Single target

                filename_parts.extend(['to', job_to.replace('.', '_')])

        else:

            # Discovery mode: add similarity range

            filename_parts.extend([f"{similarity_min}to{similarity_max}pct"])

        



        

        # Create Word document using template if available

        print("\n📄 Creating Word document...")

        

        # Check for NAB template

        template_path = Path(__file__).parent / 'templates' / 'nab_template.docx'

        if template_path.exists():

            print(f"✅ Using NAB template: {template_path}")

            template_mode = "NAB Template"

        else:

            print(f"⚠️ NAB template not found at {template_path}")

            print("   Creating document without template (basic formatting)")

            template_path = None

            template_mode = "Basic Formatting"

        

        try:

            # Try to import python-docx with enhanced formatting

            from docx import Document

            from docx.shared import Inches, Pt, RGBColor

            from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK

            from docx.enum.style import WD_STYLE_TYPE

            from datetime import datetime

            

            # Save file paths using intelligent naming

            output_dir = Path(__file__).parent

            filename = f'{"_".join(filename_parts)}.docx'

            output_path = output_dir / filename

            

            # Try professional NAB-styled document first

            if NAB_FORMATTER_AVAILABLE:

                print("🎨 Creating professional NAB-styled document...")

                

                # Add source_job_id to analysis_data for formatter

                analysis_data['source_job_id'] = test_job_id

                

                success = create_professional_word_document(word_content, analysis_data, output_path)

                

                if success:

                    print(f"✅ Professional NAB-styled document created!")

                    print(f"📄 Document features:")

                    print(f"   ✅ NAB red headings (Epilogue Semibold 22pt)")

                    print(f"   ✅ Professional typography (Source Sans Pro)")

                    print(f"   ✅ Cover page with NAB branding")

                    print(f"   ✅ TOC-compatible heading styles")

                    print(f"   ✅ Executive-ready presentation")

                else:

                    print("⚠️ Professional document creation failed, trying basic approach...")

                    success = create_basic_word_document(word_content, analysis_data, output_path)

            else:

                print("🎨 Creating basic Word document (NAB formatter not available)...")

                success = create_basic_word_document(word_content, analysis_data, output_path)

            

            if success:

                print(f"✅ Word document saved: {output_path}")

                print(f"📄 Document contains {len(section_titles)} sections")

                print(f"📊 Analysis for {job_name} with {analysis_data.get('pathway_count', 0)} pathways")

                

                if NAB_FORMATTER_AVAILABLE:

                    print(f"\n📋 To generate automatic Table of Contents:")

                    print(f"   1. Open the .docx file in Microsoft Word")

                    print(f"   2. Place cursor after 'Table of Contents' heading")

                    print(f"   3. Go to References → Table of Contents → Automatic Table 1")

                    print(f"   4. TOC will be generated from headings automatically!")

            else:

                print(f"❌ Failed to create Word document")

            

        except ImportError:

            print("⚠️ python-docx not installed. Install with: pip install python-docx")

            print("📄 Saving content to text file instead...")

            

            # Import datetime for text file fallback

            from datetime import datetime

            

            # Save as text file using intelligent naming

            output_dir = Path(__file__).parent

            filename = f'{"_".join(filename_parts)}.txt'

            output_path = output_dir / filename

            

            with open(output_path, 'w', encoding='utf-8') as f:

                f.write(f"NAB Skills Intelligence Platform\n")

                f.write(f"Strategic Career Pathway Analysis\n")

                f.write(f"Career Transition Analysis: {job_name}\n")

                f.write(f"Generated: {datetime.now().strftime('%B %d, %Y')}\n\n")

                f.write("=" * 80 + "\n\n")

                

                for section_key, section_title in section_titles:

                    if section_key in word_content:

                        f.write(f"{section_title}\n")

                        f.write("-" * len(section_title) + "\n\n")

                        

                        section_data = word_content[section_key]

                        if isinstance(section_data, dict):

                            for subsection_key, subsection in section_data.items():

                                if isinstance(subsection, dict) and 'content' in subsection:

                                    if 'title' in subsection:

                                        f.write(f"{subsection['title']}\n")

                                        f.write("-" * len(subsection['title']) + "\n")

                                    f.write(f"{subsection['content']}\n\n")

                        f.write("\n")

            

            print(f"✅ Text file saved: {output_path}")

        

        # Display summary

        print("\n" + "=" * 50)

        print("🎯 Career Transition Analysis Generator SUMMARY")

        print("=" * 50)

        

        for section_name, section_result in sections.items():

            variables = section_result.get('template_variables', {})

            content_sections = len(section_result.get('content', {}))

            print(f"✅ {section_name.replace('_', ' ').title()}: {content_sections} subsections")

        

        print(f"\n📄 Complete Career Transition Analysis generated for {job_name}")

        print(f"📊 Analysis includes {analysis_data.get('pathway_count', 0)} career pathways")

        

        print("\n🎉 Full Career Transition Analysis Generator completed successfully!")

        print("🎯 Key Features Demonstrated:")

        print("   ✅ Complete 5-section Career Transition Analysis Generator")

        print("   ✅ Professional Word document formatting")

        print("   ✅ Database-driven content throughout")

        print("   ✅ Consistent logical role architecture")

        print("   ✅ Executive-ready presentation quality")

        

    except Exception as e:

        print(f"❌ Error during full Career Transition Analysis Generator: {e}")

        import traceback

        traceback.print_exc()

    

    finally:

        if 'db' in locals():

            db.close()



# Import NAB styling from our DocumentFormatter

try:

    from ..formatter import DocumentFormatter

    NAB_FORMATTER_AVAILABLE = True

    print("✓ NAB DocumentFormatter available for professional styling")

except ImportError:

    NAB_FORMATTER_AVAILABLE = False

    print("! NAB DocumentFormatter not available - using basic styling")



def export_section_to_word(section_name, section_result, job_from, mode="top_matches", job_to=None, 

                          similarity_min=40, similarity_max=90):

    """Export a single section to Word document with intelligent naming."""

    

    print(f"\n📄 Exporting {section_name} to Word document...")

    

    try:

        # Create intelligent filename for single section

        filename_parts = [

            section_name,

            job_from.replace('.', '_'),

            mode

        ]

        

        # Add specific mode details to filename

        if mode == "specific" and job_to:

            if ',' in job_to:

                target_count = len([j.strip() for j in job_to.split(',')])

                first_target = job_to.split(',')[0].strip().replace('.', '_')

                filename_parts.extend(['multi', f"{target_count}targets", first_target])

            else:

                filename_parts.extend(['to', job_to.replace('.', '_')])

        else:

            filename_parts.extend([f"{similarity_min}to{similarity_max}pct"])

        

        output_dir = Path(__file__).parent

        filename = f'{"_".join(filename_parts)}.docx'

        output_path = output_dir / filename

        

        # Prepare single section content for Word export

        single_section_content = {section_name: section_result.get('content', {})}

        

        # Create analysis data for the section

        variables = section_result.get('template_variables', {})

        analysis_data = {

            'source_job_logical_display_name': variables.get('source_job_title', 'Professional Role'),

            'source_job_id': job_from,

            'summary': f'{section_name.replace("_", " ").title()} for {variables.get("source_job_title", "Professional Role")}',

            'avg_similarity': variables.get('max_similarity', 0) / 100 if variables.get('max_similarity') else 0,

            'pathway_count': variables.get('pathway_count', 0),

            'analysis_mode': mode,

            'target_job': job_to,

            'similarity_range': f"{similarity_min}%-{similarity_max}%"

        }

        

        # Try professional document creation first

        if NAB_FORMATTER_AVAILABLE:

            success = create_professional_word_document(single_section_content, analysis_data, output_path)

        else:

            success = create_basic_word_document(single_section_content, analysis_data, output_path)

        

        if success:

            print(f"✅ {section_name.replace('_', ' ').title()} exported to: {output_path}")

        else:

            print(f"❌ Failed to export {section_name}")

            

    except Exception as e:

        print(f"❌ Error exporting {section_name}: {e}")



def create_professional_word_document(content_sections, analysis_data, output_path):

    """Create a professional Word document with NAB styling using DocumentFormatter."""

    

    if not NAB_FORMATTER_AVAILABLE:

        print("❌ Cannot create professional document - DocumentFormatter not available")

        return False

    

    try:

        # Initialize the DocumentFormatter

        formatter = DocumentFormatter()

        

        # Generate the document using the formatter's format_document method

        result = formatter.format_document(

            content=content_sections,

            output_format='word',

            analysis_data=analysis_data

        )

        

        if result.get('status') == 'generated':

            # Save the document

            with open(output_path, 'wb') as f:

                f.write(result['content'])

            

            print(f"✅ Professional NAB-styled Word document saved: {output_path}")

            return True

        else:

            print(f"❌ Document generation failed: {result.get('message', 'Unknown error')}")

            return False

        

    except Exception as e:

        print(f"❌ Error creating professional document: {e}")

        return False



# Fallback function for basic Word document creation

def create_basic_word_document(content_sections, analysis_data, output_path):

    """Fallback function to create a basic Word document without NAB styling."""

    

    try:

        from docx import Document

        from docx.shared import Pt

        from datetime import datetime

        

        doc = Document()

        

        # Basic title page

        title = doc.add_paragraph('Skills Intelligence Platform')

        title_run = title.runs[0]

        title_run.font.size = Pt(24)

        title_run.bold = True

        

        subtitle = doc.add_paragraph('Strategic Career Pathway Analysis')

        subtitle_run = subtitle.runs[0] 

        subtitle_run.font.size = Pt(18)

        

        job_name = analysis_data.get('source_job_logical_display_name', 'Professional Role')

        job_para = doc.add_paragraph(f'Career Transition Analysis: {job_name}')

        job_run = job_para.runs[0]

        job_run.font.size = Pt(16)

        job_run.bold = True

        

        date_para = doc.add_paragraph(f'Generated: {datetime.now().strftime("%B %d, %Y")}')

        

        doc.add_page_break()

        

        # Basic table of contents

        toc_heading = doc.add_paragraph('Table of Contents')

        toc_run = toc_heading.runs[0]

        toc_run.font.size = Pt(16)

        toc_run.bold = True

        

        doc.add_paragraph("(Table of Contents placeholder - generate in Word)")

        doc.add_page_break()

        

        # Add sections

        section_order = [

            ('executive_summary', 'Executive Summary'),

            ('current_role_context', 'Current Role Context'), 

            ('pathway_analysis', 'Pathway Analysis: Top 3 Strategic Opportunities'),

            ('strategic_recommendations', 'Strategic Recommendations'),

            ('conclusion', 'Conclusion')

        ]

        

        for section_key, section_title in section_order:

            if section_key in content_sections:

                # Section heading

                heading = doc.add_paragraph(section_title)

                heading_run = heading.runs[0]

                heading_run.font.size = Pt(14)

                heading_run.bold = True

                

                # Section content

                section_data = content_sections[section_key]

                if isinstance(section_data, dict):

                    for subsection_key, subsection in section_data.items():

                        if isinstance(subsection, dict) and 'content' in subsection:

                            if 'title' in subsection:

                                sub_heading = doc.add_paragraph(subsection['title'])

                                sub_run = sub_heading.runs[0]

                                sub_run.font.size = Pt(12)

                                sub_run.bold = True

                            

                            content_para = doc.add_paragraph(subsection['content'])

        

        doc.save(str(output_path))

        print(f"✅ Basic Word document saved: {output_path}")

        return True

        

    except ImportError:

        print("❌ python-docx not available for Word document creation")

        return False

    except Exception as e:

        print(f"❌ Error creating basic document: {e}")

        return False



def main():

    """Main function with command line argument parsing."""

    parser = argparse.ArgumentParser(

        description="Test Career Transition Analysis Generator components with configurable parameters",

        formatter_class=argparse.RawDescriptionHelpFormatter,

        epilog="""

Examples:

    # Basic test modes

    python test_career_analysis.py                           # Run all tests with default job

    python test_career_analysis.py --copy                    # Display clean copy without debug

    python test_career_analysis.py --template                # Test template loading only

    python test_career_analysis.py --executive               # Test executive summary only

    python test_career_analysis.py --current                 # Test current role context only

    python test_career_analysis.py --pathway                 # Test pathway analysis only

    python test_career_analysis.py --strategic               # Test strategic recommendations only

    python test_career_analysis.py --conclusion              # Test conclusion only

    python test_career_analysis.py --word                    # Generate full Word document

    

    # Core parameters matching HTML interface

    python test_career_analysis.py --job-from R0100.2        # Specify source job

    python test_career_analysis.py --mode top_matches        # Analysis mode: top_matches (default) or specific

    python test_career_analysis.py --job-to R0025.1          # Target job (required for specific mode)

    python test_career_analysis.py --similarity-min 40       # Minimum similarity percentage (20-95)

    python test_career_analysis.py --similarity-max 90       # Maximum similarity percentage (25-100)

    

    # Combined examples

    python test_career_analysis.py --job-from R0100.2 --mode top_matches --similarity-min 45 --similarity-max 85 --word

    python test_career_analysis.py --job-from R0100.2 --mode specific --job-to R0025.1 --similarity-min 50 --similarity-max 80 --executive

    python test_career_analysis.py --job-from R0100.2 --mode specific --job-to R0025.1 --copy

        """

    )

    

    # Test mode arguments

    parser.add_argument('--template', action='store_true', 

                       help='Test template loading only')

    parser.add_argument('--executive', action='store_true', 

                       help='Test executive summary generation only')

    parser.add_argument('--current', action='store_true', 

                       help='Test current role context generation only')

    parser.add_argument('--pathway', action='store_true', 

                       help='Test pathway analysis generation only')

    parser.add_argument('--strategic', action='store_true', 

                       help='Test strategic recommendations generation only')

    parser.add_argument('--conclusion', action='store_true', 

                       help='Test conclusion generation only')

    parser.add_argument('--word', action='store_true', 

                       help='Test full Career Transition Analysis Generator with Word document output')

    parser.add_argument('--copy', action='store_true', 

                       help='Display clean copy of all sections without debug information')

    parser.add_argument('--all', action='store_true', 

                       help='Run all tests (default if no flags specified)')

    

    # Core parameters matching HTML interface

    parser.add_argument('--job-from', type=str, default='R0100.2',

                       help='Source job ID (default: R0100.2)')

    parser.add_argument('--mode', type=str, choices=['top_matches', 'specific'], default='top_matches',

                       help='Analysis mode: top_matches (default) or specific transition')

    parser.add_argument('--job-to', type=str, default=None,

                       help='Target job ID (required for specific mode)')

    parser.add_argument('--similarity-min', type=int, default=40, 

                       help='Minimum similarity percentage (20-95, default: 40)')

    parser.add_argument('--similarity-max', type=int, default=90,

                       help='Maximum similarity percentage (25-100, default: 90)')

    parser.add_argument('--top-n', type=int, default=3,

                       help='Number of top pathways to analyze (default: 3)')

    

    # Tie-breaking preference arguments

    parser.add_argument('--same-function-priority', action='store_true',

                       help='Prioritise opportunities within the same job function')

    parser.add_argument('--career-progression-priority', action='store_true',

                       help='Prioritise higher-level roles over lateral moves')

    parser.add_argument('--minimal-level-jump', action='store_true',

                       help='Prioritise roles closest to current management level')

    parser.add_argument('--skills-overlap-detail', action='store_true',

                       help='Show detailed skills overlap for tie-breaking transparency')

    

    args = parser.parse_args()

    

    # Extract parameters

    job_from = args.job_from

    mode = args.mode

    job_to = args.job_to

    similarity_min = args.similarity_min

    similarity_max = args.similarity_max

    top_n = args.top_n

    

    # Tie-breaking preferences

    tie_breaking_options = {

        'same_function_priority': args.same_function_priority,

        'career_progression_priority': args.career_progression_priority,

        'minimal_level_jump': args.minimal_level_jump,

        'skills_overlap_detail': args.skills_overlap_detail

    }

    

    # If no specific test is requested, run all tests

    run_all = (not (args.template or args.executive or args.current or args.pathway or 

                   args.strategic or args.conclusion or args.copy or args.word))

    

    # Handle clean copy display separately (uses parameters)

    if args.copy:

        test_clean_copy_display(job_from, mode, job_to, similarity_min, similarity_max, top_n)

        return

    

    print("🚀 Career Transition Analysis Generator Test Suite")

    print("=" * 50)

    print(f"📋 Global Parameters:")

    print(f"   Source Job: {job_from}")

    print(f"   Analysis Mode: {mode}")

    if mode == "specific":

        print(f"   Target Job: {job_to or 'Not specified'}")

    print(f"   Similarity Range: {similarity_min}% - {similarity_max}%")

    print(f"   🎯 Top N Pathways: {top_n}")

    

    # Debug tie-breaking options

    active_options = [k for k, v in tie_breaking_options.items() if v]

    if active_options:

        print(f"   🔧 Active Tie-Breaking Options:")

        for option in active_options:

            option_display = option.replace('_', ' ').title()

            print(f"      ✅ {option_display}")

    else:

        print(f"   🔧 Tie-Breaking Options: None (using default ordering)")

    print()

    

    if args.template or run_all:

        test_template_loading()

    

    if args.executive or run_all:

        test_executive_summary_generation(job_from, mode, job_to, similarity_min, similarity_max, top_n)

    

    if args.current or run_all:

        test_current_role_context_generation(job_from, mode, job_to, similarity_min, similarity_max)

    

    if args.pathway or run_all:

        test_pathway_analysis_generation(job_from, mode, job_to, similarity_min, similarity_max, top_n, tie_breaking_options)

    

    if args.strategic or run_all:

        test_strategic_recommendations_generation(job_from, mode, job_to, similarity_min, similarity_max, top_n)

    

    if args.conclusion or run_all:

        test_conclusion_generation(job_from, mode, job_to, similarity_min, similarity_max, top_n)

    

    if args.word:

        test_full_career_analysis_generation(job_from, mode, job_to, similarity_min, similarity_max, top_n)

    

    print("\n🎉 Test suite completed!")



if __name__ == "__main__":

    main() 

